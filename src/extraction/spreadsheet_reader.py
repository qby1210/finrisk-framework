"""CSV/XLSX 表格文件读取器。"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook  # type: ignore[import-untyped]

from domain.common import ExtractionMethod
from domain.file_models import ExtractedTable
from extraction.types import DocumentDraft, ElementDraft, PageDraft

_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030")


def read_spreadsheet(path: Path) -> DocumentDraft:
    """读取 CSV 或 XLSX；每个工作表作为一个逻辑页，并保留原始单元格行列。"""

    if path.suffix.lower() == ".csv":
        return _read_csv(path)
    if path.suffix.lower() == ".xlsx":
        return _read_xlsx(path)
    raise ValueError(f"不支持的表格格式：{path.suffix}")


def _read_csv(path: Path) -> DocumentDraft:
    text: str | None = None
    used_encoding: str | None = None
    for encoding in _ENCODINGS:
        try:
            text = path.read_text(encoding=encoding)
            used_encoding = encoding
            break
        except UnicodeDecodeError:
            continue
    if text is None or used_encoding is None:
        raise UnicodeError(f"无法识别 CSV 文件编码：{path.name}")
    rows = [[cell.strip() for cell in row] for row in csv.reader(text.splitlines())]
    rows = [row for row in rows if any(row)]
    warning = (
        [] if used_encoding.startswith("utf-8") else [f"CSV 使用兼容编码读取：{used_encoding}"]
    )
    return DocumentDraft(
        pages=[_sheet_page(path.stem, rows, warning)],
        tables=[
            ExtractedTable(
                page_number=1,
                page_numbers=[1],
                title=path.stem,
                rows=rows,
                confidence=1.0,
            )
        ]
        if rows
        else [],
        warnings=warning,
    )


def _read_xlsx(path: Path) -> DocumentDraft:
    workbook = load_workbook(path, read_only=True, data_only=True)
    pages: list[PageDraft] = []
    tables: list[ExtractedTable] = []
    warnings = ["XLSX 工作表按逻辑页处理，不等同于打印物理页码"]
    try:
        for page_number, worksheet in enumerate(workbook.worksheets, start=1):
            rows = [
                ["" if cell is None else str(cell).strip() for cell in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            rows = [row for row in rows if any(row)]
            pages.append(_sheet_page(worksheet.title, rows, []))
            if rows:
                tables.append(
                    ExtractedTable(
                        page_number=page_number,
                        page_numbers=[page_number],
                        title=worksheet.title,
                        rows=rows,
                        confidence=1.0,
                    )
                )
    finally:
        workbook.close()
    return DocumentDraft(pages=pages, tables=tables, warnings=warnings)


def _sheet_page(title: str, rows: list[list[str]], warnings: list[str]) -> PageDraft:
    blocks = [
        ElementDraft(
            text=f"【工作表：{title}】",
            confidence=None,
            polygon=None,
            extraction_method=ExtractionMethod.SPREADSHEET,
        )
    ]
    blocks.extend(
        ElementDraft(
            text=" | ".join(cell for cell in row if cell),
            confidence=None,
            polygon=None,
            extraction_method=ExtractionMethod.SPREADSHEET,
        )
        for row in rows
        if any(row)
    )
    return PageDraft(
        width=None,
        height=None,
        blocks=blocks,
        words=[],
        used_ocr=False,
        average_confidence=None,
        quality_score=1.0 if rows else 0.0,
        extraction_method=ExtractionMethod.SPREADSHEET,
        warnings=warnings,
    )
